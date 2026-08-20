from pathlib import Path
from llama_index.core import Document
from llama_index.readers.file import PDFReader
from docx import Document as DocxDocument


class EnterpriseLoader:

    def load_pdf(self, file_path):
        reader = PDFReader()
        docs = reader.load_data(file=file_path)

        for doc in docs:
            doc.metadata["source"] = Path(file_path).name
            doc.metadata["type"] = "pdf"

        return docs

    def load_docx(self, file_path):

        docx = DocxDocument(file_path)

        text = "\n".join(
            [p.text for p in docx.paragraphs if p.text.strip()]
        )

        doc = Document(
            text=text,
            metadata={
                "source": Path(file_path).name,
                "type": "docx"
            }
        )

        return [doc]

    def load_txt(self, file_path):

        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read()

        doc = Document(
            text=text,
            metadata={
                "source": Path(file_path).name,
                "type": "txt"
            }
        )

        return [doc]